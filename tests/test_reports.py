"""SQL reports agree with pandas; the Excel workbook has live formulas that recalculate correctly."""

import math
import shutil
import subprocess

import numpy as np
import openpyxl
import pandas as pd
import pytest

from wfm.capacity import interval_requirement
from wfm.egypt_calendar import week_start
from wfm.excel_report import agent_minutes, build_workbook
from wfm.forecast import interval_forecast, learn_profiles
from wfm.simulate import simulate
from wfm.sql_reports import REPORTS, build_database, run_report

PARAMS = {"queue": "AR_TECH", "date": "2026-09-10", "start": "2026-08-01", "end": "2026-09-30"}


@pytest.fixture(scope="module")
def sim():
    return simulate("2026-07-01", "2026-09-30", seed=4)


@pytest.fixture(scope="module")
def conn(sim, rta_day):
    return build_database(sim["intervals"], sim["calendar"], sim["events"], rta_day)


def test_every_report_runs(conn):
    for key in REPORTS:
        assert len(run_report(conn, key, **PARAMS)) > 0, key


def test_missing_parameter_is_reported(conn):
    with pytest.raises(ValueError, match="needs parameters"):
        run_report(conn, "interval_report", queue="AR_TECH")


def test_interval_report_matches_pandas(sim, conn):
    iv = sim["intervals"]
    day = iv[(iv["queue"] == "AR_TECH") & (iv["date"] == PARAMS["date"])].sort_values("interval")
    rep = run_report(conn, "interval_report", **PARAMS)
    assert len(rep) == 48
    assert rep["offered"].tolist() == day["offered"].tolist()
    np.testing.assert_allclose(rep["sl_pct"], (100 * day["answered_in_sl"] / day["offered"]).round(1).to_numpy())


def test_daily_kpis_window_functions(sim, conn):
    iv = sim["intervals"]
    q = iv[iv["queue"] == "AR_TECH"]
    daily = q.groupby("date").agg(offered=("offered", "sum"), in_sl=("answered_in_sl", "sum"))
    rolling = (100 * daily["in_sl"].rolling(7, min_periods=1).sum() / daily["offered"].rolling(7, min_periods=1).sum()).round(1)
    rep = run_report(conn, "daily_kpis", **PARAMS).set_index("date")
    expected = rolling[(rolling.index >= PARAMS["start"]) & (rolling.index <= PARAMS["end"])]
    np.testing.assert_allclose(rep["sl_7day_pct"].to_numpy(), expected.to_numpy())
    assert rep.index.min() == PARAMS["start"]  # the filter is applied after the windows


def test_adherence_report_matches_agent_minutes(conn, rta_day):
    rep = run_report(conn, "adherence_report").set_index("agent_id")
    mins = agent_minutes(rta_day).set_index("agent_id")
    assert set(rep.index) == set(mins.index)
    np.testing.assert_array_equal(rep.loc[mins.index, "out_of_adherence_min"], mins["out_of_adherence_min"])


@pytest.fixture(scope="module")
def workbook(tmp_path_factory, sim, rta_day):
    clean = sim["intervals"].assign(offered_clean=sim["intervals"]["offered"].astype(float), is_outage=False)
    profiles = learn_profiles(clean, sim["calendar"], pd.Timestamp("2026-09-30"))
    dates = pd.date_range("2026-10-04", periods=28, freq="D")
    daily_fc = pd.DataFrame([(d, q, 5000.0) for d in dates for q in ("AR_TECH",)], columns=["date", "queue", "forecast"])
    req = interval_requirement(interval_forecast(daily_fc, sim["calendar"], profiles))
    req["week_start"] = week_start(req["date"])
    path = tmp_path_factory.mktemp("xl") / "wfm.xlsx"
    return build_workbook(path, rta_day, agent_minutes(rta_day), sim["intervals"], sim["calendar"], req)


def test_workbook_sheets_and_formulas(workbook):
    wb = openpyxl.load_workbook(workbook)
    assert wb.sheetnames == ["README", "Erlang C", "Interval report", "Shrinkage", "Capacity plan", "Adherence"]
    erlang = wb["Erlang C"]
    assert erlang["B12"].value.startswith("=INDEX(")
    assert "POISSON.DIST" in erlang["B15"].value
    assert wb["Shrinkage"]["C12"].value.startswith("=1-EXP(SUMPRODUCT(LN(")
    values = openpyxl.load_workbook(workbook, data_only=True)
    assert values["Erlang C"]["B12"].value == 14  # textbook: 100 calls, 180s AHT, 80/20 -> 14 agents
    assert math.isclose(values["Erlang C"]["B13"].value, 0.8884, abs_tol=1e-4)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_workbook_recalculates_in_libreoffice(workbook, tmp_path):
    profile = tmp_path / "profile" / "user"
    profile.mkdir(parents=True)
    (profile / "registrymodifications.xcu").write_text(
        '<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
        '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
        "<value>0</value></prop></item></oor:items>"
    )
    subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{tmp_path / 'profile'}", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "out"), str(workbook)],
        check=True, capture_output=True, timeout=180,
    )
    cached = openpyxl.load_workbook(workbook, data_only=True)
    formulas = openpyxl.load_workbook(workbook)
    recalculated = openpyxl.load_workbook(tmp_path / "out" / workbook.name, data_only=True)
    checked = 0
    for ws in formulas.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    expected, got = cached[ws.title][cell.coordinate].value, recalculated[ws.title][cell.coordinate].value
                    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
                        assert math.isclose(float(got), float(expected), rel_tol=1e-6, abs_tol=1e-6), (ws.title, cell.coordinate)
                    else:
                        assert expected == got or (expected in (None, "") and got in (None, "")), (ws.title, cell.coordinate)
                    checked += 1
    assert checked > 500
